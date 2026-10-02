"""Bulk product import from Excel (.xlsx) or CSV, used by the Products admin.

Flow: build_template() -> the editor fills it -> read_file() -> validate_rows()
shows a preview -> apply_rows() saves everything in one transaction.
Imported products have no main image, so the website hides them until one is
added in the admin.
"""
import csv
import io
import zipfile
from collections import defaultdict

from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.db import transaction
from django.utils.text import slugify
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from categories.models import Category
from companies.models import Company, CompanyLine
from .models import Product

MAX_ROWS = 2000
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
SEO_TITLE_TIP = 60
SEO_DESCRIPTION_TIP = 160

# (key, header in the template, required)
COLUMNS = [
    ('name', 'Name', True),
    ('brand', 'Brand', True),
    ('line', 'Product line', False),
    ('category', 'Category', True),
    ('description', 'Description', True),
    ('sizes', 'Sizes', False),
    ('available', 'Available', False),
    ('seo_title', 'SEO title', False),
    ('seo_description', 'SEO description', False),
]
COLUMN_WIDTHS = [32, 22, 24, 28, 60, 22, 12, 40, 60]

# Other spellings accepted in the header row. "YouTube link" is accepted but
# left out of the template, since videos are usually added in the admin.
HEADER_ALIASES = {
    'name': 'name', 'product name': 'name',
    'brand': 'brand', 'company': 'brand',
    'product line': 'line', 'line': 'line', 'company line': 'line',
    'category': 'category',
    'description': 'description',
    'sizes': 'sizes', 'size': 'sizes',
    'available': 'available', 'is available': 'available',
    'seo title': 'seo_title', 'meta title': 'seo_title',
    'seo description': 'seo_description', 'meta description': 'seo_description',
    'youtube link': 'youtube', 'youtube url': 'youtube', 'youtube': 'youtube',
}
REQUIRED_KEYS = [key for key, _, required in COLUMNS if required]
HEADER_LABELS = {key: label for key, label, _ in COLUMNS} | {'youtube': 'YouTube link'}

YES_VALUES = {'yes', 'y', 'true', '1'}
NO_VALUES = {'no', 'n', 'false', '0'}


class ImportAborted(Exception):
    """Something changed between the preview and the confirm step."""


def _clean_header(value):
    return ' '.join(str(value or '').replace('_', ' ').split()).lower()


def _clean_cell(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------

def build_template():
    """The Excel file editors fill in. Brand, Product line and Category have
    dropdowns of the names currently in the admin, refreshed on every download."""
    companies = list(Company.objects.order_by('name'))
    lines = list(CompanyLine.objects.select_related('company').order_by('company__name', 'name'))
    categories = list(
        Category.objects.select_related('company', 'company_line').order_by('company__name', 'name')
    )

    workbook = Workbook()
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='11100E')

    # Sheet 1: the products to fill in
    sheet = workbook.active
    sheet.title = 'Products'
    for index, (_, label, required) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.font = header_font
        cell.fill = PatternFill('solid', fgColor='70443D') if required else header_fill
        cell.alignment = Alignment(vertical='center')
        sheet.column_dimensions[cell.column_letter].width = COLUMN_WIDTHS[index - 1]
    sheet.row_dimensions[1].height = 22
    sheet.freeze_panes = 'A2'

    # Hidden sheet holding the dropdown lists
    lists = workbook.create_sheet('_lists')
    lists.sheet_state = 'hidden'
    dropdowns = {
        'brand': sorted({company.name for company in companies}),
        'line': sorted({line.name for line in lines}),
        'category': sorted({category.name for category in categories}),
        'available': ['Yes', 'No'],
    }
    column_of = {key: index for index, (key, _, _) in enumerate(COLUMNS, start=1)}
    last_row = MAX_ROWS + 1
    for list_column, (key, values) in enumerate(dropdowns.items(), start=1):
        if not values:
            continue
        for row, value in enumerate(values, start=1):
            lists.cell(row=row, column=list_column, value=value)
        letter = lists.cell(row=1, column=list_column).column_letter
        validation = DataValidation(
            type='list',
            formula1=f"'_lists'!${letter}$1:${letter}${len(values)}",
            allow_blank=True,
            showErrorMessage=True,
            errorStyle='warning',
            errorTitle='Name not in the admin',
            error='Pick a name from the list. Names must match the admin exactly.',
        )
        target = sheet.cell(row=2, column=column_of[key]).column_letter
        validation.add(f'{target}2:{target}{last_row}')
        sheet.add_data_validation(validation)

    # Sheet 2: which lines and categories belong to which brand
    allowed = workbook.create_sheet('Allowed names')
    for index, label in enumerate(['Brand', 'Product line', 'Category'], start=1):
        cell = allowed.cell(row=1, column=index, value=label)
        cell.font = header_font
        cell.fill = header_fill
        allowed.column_dimensions[cell.column_letter].width = 32
    allowed.freeze_panes = 'A2'
    allowed_rows = [
        (
            category.company.name if category.company else '(any brand)',
            category.company_line.name if category.company_line else '',
            category.name,
        )
        for category in categories
    ]
    used_lines = {(row[0], row[1]) for row in allowed_rows}
    allowed_rows += [
        (line.company.name, line.name, '')
        for line in lines
        if (line.company.name, line.name) not in used_lines
    ]
    for row in sorted(allowed_rows):
        allowed.append(row)

    # Sheet 3: short instructions
    guide = workbook.create_sheet('How to fill')
    guide.column_dimensions['A'].width = 110
    for line in [
        'HOW TO FILL THIS FILE',
        '',
        '1. Fill the "Products" sheet: one row per product. Do not change the first row.',
        '2. Columns in brown (Name, Brand, Category, Description) are required.',
        '3. Brand, Product line and Category must match the admin exactly - pick them from the dropdowns.',
        '   The "Allowed names" sheet shows which product lines and categories belong to each brand.',
        '4. A category must belong to the same brand as the product.',
        '   New categories must be created in admin -> Categories BEFORE importing.',
        '5. Sizes: separate with commas, e.g.  S, M, L, XL',
        '6. Available: Yes or No (empty means Yes).',
        '7. SEO title (best under 60 characters) and SEO description (120-155) are optional.',
        '   " | Muttrah Pharmacy" is added to the SEO title automatically.',
        '',
        'AFTER IMPORTING',
        'Products are saved without images and stay hidden on the website.',
        'Open each product in admin -> Products, upload its main image and Save: it then appears on the website.',
        'Tip: in admin -> Products use the filter "Has main image: No" to see which still need a photo.',
        '',
        'A product that already exists (same Name and Brand) is updated instead of added again.',
        'Empty optional cells never erase what is already saved for that product.',
    ]:
        guide.append([line])
    guide['A1'].font = Font(bold=True, size=13)
    guide['A14'].font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Reading the uploaded file
# ---------------------------------------------------------------------------

def _read_csv(uploaded):
    raw = uploaded.read()
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError('The CSV file could not be read. Save it as "CSV UTF-8" and try again.')
    return [(index, row) for index, row in enumerate(csv.reader(io.StringIO(text)), start=1)]


def _read_xlsx(uploaded):
    try:
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError):
        raise ValueError('The file could not be opened. Please upload an Excel (.xlsx) file.')
    try:
        sheet = workbook['Products'] if 'Products' in workbook.sheetnames else workbook.worksheets[0]
        return [
            (index, list(row))
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)
        ]
    finally:
        workbook.close()


def read_file(uploaded):
    """Returns (rows, error). Each row is a dict of column key -> text, plus
    "row", the row number in the file (as Excel shows it)."""
    name = (uploaded.name or '').lower()
    try:
        raw_rows = _read_csv(uploaded) if name.endswith('.csv') else _read_xlsx(uploaded)
    except ValueError as error:
        return [], str(error)

    # The header is the first row that has anything in it
    header_index = next(
        (position for position, (_, cells) in enumerate(raw_rows) if any(_clean_cell(c) for c in cells)),
        None,
    )
    if header_index is None:
        return [], 'The file is empty.'

    columns = {}
    unknown = []
    for position, cell in enumerate(raw_rows[header_index][1]):
        header = _clean_header(cell)
        if not header:
            continue
        key = HEADER_ALIASES.get(header)
        if key and key not in columns:
            columns[key] = position
        elif not key:
            unknown.append(str(cell).strip())

    missing = [HEADER_LABELS[key] for key in REQUIRED_KEYS if key not in columns]
    if missing:
        return [], (
            f'Missing column(s): {", ".join(missing)}. '
            'Please use the downloaded template and do not change its first row.'
        )

    rows = []
    for row_number, cells in raw_rows[header_index + 1:]:
        row = {
            key: _clean_cell(cells[position]) if position < len(cells) else ''
            for key, position in columns.items()
        }
        if not any(row.values()):
            continue
        row['row'] = row_number
        rows.append(row)

    if not rows:
        return [], 'No products found in the file. Fill in at least one row under the headers.'
    if len(rows) > MAX_ROWS:
        return [], f'The file has {len(rows)} products. Please import at most {MAX_ROWS} at a time.'
    if unknown:
        rows[0].setdefault('_notes', []).append(
            f'These columns were ignored: {", ".join(unknown)}.'
        )
    return rows, ''


# ---------------------------------------------------------------------------
# Checking every row
# ---------------------------------------------------------------------------

def _parse_available(value):
    text = value.strip().lower()
    if not text or text in YES_VALUES:
        return True
    if text in NO_VALUES:
        return False
    return None


def validate_rows(rows):
    """Checks every row against the admin data. Returns a list of results
    (plain data, safe to keep in the session until the editor confirms)."""
    companies = defaultdict(list)
    for company in Company.objects.all():
        companies[company.name.strip().lower()].append(company)

    lines = defaultdict(list)
    for line in CompanyLine.objects.all():
        lines[(line.company_id, line.name.strip().lower())].append(line)

    categories = defaultdict(list)
    for category in Category.objects.select_related('company', 'company_line'):
        categories[category.name.strip().lower()].append(category)

    existing = defaultdict(list)
    for product_id, company_id, product_name in Product.objects.values_list('id', 'company_id', 'name'):
        existing[(company_id, product_name.strip().lower())].append(product_id)

    url_validator = URLValidator()
    seen = {}
    results = []

    for row in rows:
        errors = []
        warnings = list(row.get('_notes', []))
        name = row.get('name', '')
        brand_text = row.get('brand', '')
        line_text = row.get('line', '')
        category_text = row.get('category', '')
        description = row.get('description', '')
        sizes = row.get('sizes', '')
        seo_title = row.get('seo_title', '')
        seo_description = row.get('seo_description', '')
        youtube = row.get('youtube', '')

        company = line = category = None

        # Name
        if not name:
            errors.append('Name is empty.')
        elif len(name) > 100:
            errors.append(f'Name is too long ({len(name)} characters, maximum 100).')

        # Brand
        if not brand_text:
            errors.append('Brand is empty.')
        else:
            matches = companies.get(brand_text.lower(), [])
            if len(matches) == 1:
                company = matches[0]
            elif matches:
                errors.append(f'More than one brand is named "{brand_text}" in the admin.')
            else:
                errors.append(f'Brand "{brand_text}" not found. Check the spelling in admin → Companies.')

        # Product line (optional) must belong to the brand
        if line_text and company:
            matches = lines.get((company.id, line_text.lower()), [])
            if matches:
                line = matches[0]
            else:
                errors.append(f'Product line "{line_text}" is not a line of {company.name}.')

        # Category must belong to the brand (or to no brand)
        if not category_text:
            errors.append('Category is empty.')
        elif company:
            named = categories.get(category_text.lower(), [])
            candidates = [c for c in named if c.company_id == company.id] or [
                c for c in named if c.company_id is None
            ]
            if line and len(candidates) > 1:
                candidates = [c for c in candidates if c.company_line_id == line.id] or [
                    c for c in candidates if c.company_line_id is None
                ]
            if len(candidates) == 1:
                category = candidates[0]
            elif len(candidates) > 1:
                errors.append(
                    f'There is more than one category "{category_text}" for {company.name}. '
                    'Fill in Product line to choose the right one.'
                )
            elif named:
                owners = ', '.join(sorted({c.company.name for c in named if c.company}))
                errors.append(f'Category "{category_text}" belongs to {owners}, not {company.name}.')
            else:
                errors.append(
                    f'Category "{category_text}" not found. Create it in admin → Categories first.'
                )

        if category and category.company_line_id:
            if line and line.id != category.company_line_id:
                errors.append(
                    f'Category "{category.name}" belongs to product line '
                    f'"{category.company_line.name}", not "{line.name}".'
                )
            elif not line:
                line = category.company_line
                warnings.append(f'Product line set to "{line.name}" (from its category).')

        # Description
        if not description:
            errors.append('Description is empty.')

        # Optional columns
        if len(sizes) > 255:
            errors.append('Sizes is too long (maximum 255 characters).')
        available = _parse_available(row.get('available', ''))
        if available is None:
            errors.append(f'Available must be Yes or No (found "{row.get("available")}").')
        if len(seo_title) > 255:
            errors.append('SEO title is too long (maximum 255 characters).')
        elif len(seo_title) > SEO_TITLE_TIP:
            warnings.append(f'SEO title has {len(seo_title)} characters; Google may cut it after {SEO_TITLE_TIP}.')
        if len(seo_description) > SEO_DESCRIPTION_TIP:
            warnings.append(
                f'SEO description has {len(seo_description)} characters; '
                f'Google may cut it after about {SEO_DESCRIPTION_TIP}.'
            )
        if youtube:
            try:
                url_validator(youtube)
            except ValidationError:
                errors.append('YouTube link is not a valid web address.')

        # New product, or an update of one that already exists?
        product_id = None
        if company and name:
            key = (company.id, name.lower())
            if key in seen:
                errors.append(f'The same product is already in row {seen[key]}.')
            else:
                seen[key] = row['row']
            matches = existing.get(key, [])
            if len(matches) == 1:
                product_id = matches[0]
            elif len(matches) > 1:
                errors.append(
                    f'More than one existing product is named "{name}" for {company.name}. '
                    'Update it by hand in the admin.'
                )

        if errors:
            status = 'error'
        else:
            status = 'update' if product_id else 'new'

        results.append({
            'row': row['row'],
            'status': status,
            'errors': errors,
            'warnings': warnings,
            'product_id': product_id,
            'name': name,
            'brand': company.name if company else brand_text,
            'line': line.name if line else line_text,
            'category': category.name if category else category_text,
            'company_id': company.id if company else None,
            'line_id': line.id if line else None,
            'category_id': category.id if category else None,
            'description': description,
            'sizes': sizes,
            'available': available,
            'available_given': bool(row.get('available', '').strip()),
            'seo_title': seo_title,
            'seo_description': seo_description,
            'youtube': youtube,
        })

    return results


def summarize(results):
    return {
        'total': len(results),
        'new': sum(1 for r in results if r['status'] == 'new'),
        'update': sum(1 for r in results if r['status'] == 'update'),
        'error': sum(1 for r in results if r['status'] == 'error'),
        'warning': sum(1 for r in results if r['warnings']),
    }


# ---------------------------------------------------------------------------
# Saving
# ---------------------------------------------------------------------------

def _unique_slug(name, taken):
    """`taken` holds every slug already in use (loaded once, then kept up to date)."""
    base = slugify(name)[:50].strip('-') or 'product'
    candidate = base
    counter = 2
    while candidate in taken:
        suffix = f'-{counter}'
        candidate = f'{base[:50 - len(suffix)].rstrip("-")}{suffix}'
        counter += 1
    taken.add(candidate)
    return candidate


@transaction.atomic
def apply_rows(results):
    """Creates and updates the products from validated results, all or nothing.
    Returns (created, updated)."""
    if any(result['status'] == 'error' for result in results):
        raise ImportAborted('The file still has rows with errors. Please fix them and upload again.')

    company_ids = {r['company_id'] for r in results}
    line_ids = {r['line_id'] for r in results if r['line_id']}
    category_ids = {r['category_id'] for r in results}
    companies = Company.objects.in_bulk(company_ids)
    lines = CompanyLine.objects.in_bulk(line_ids)
    categories = Category.objects.in_bulk(category_ids)
    if len(companies) != len(company_ids) or len(lines) != len(line_ids) or len(categories) != len(category_ids):
        raise ImportAborted(
            'A brand, product line or category was changed or deleted in the admin after the preview. '
            'Please upload the file again.'
        )

    # Loaded once so large files don't run extra queries per row
    taken_slugs = set(Product.objects.values_list('slug', flat=True))
    to_update = Product.objects.in_bulk(
        [r['product_id'] for r in results if r['status'] == 'update']
    )

    new_products = []
    updated = 0
    for result in results:
        if result['status'] == 'update':
            product = to_update.get(result['product_id'])
            if product is None:
                raise ImportAborted(
                    f'Product "{result["name"]}" was deleted after the preview. Please upload the file again.'
                )
            product.name = result['name']
            product.category = categories[result['category_id']]
            if result['line_id']:
                product.company_line = lines[result['line_id']]
            product.description = result['description']
            # Empty optional cells keep what is already saved
            if result['sizes']:
                product.size = result['sizes']
            if result['available_given']:
                product.is_available = result['available']
            if result['seo_title']:
                product.meta_title = result['seo_title']
            if result['seo_description']:
                product.meta_description = result['seo_description']
            if result['youtube']:
                product.youtube_url = result['youtube']
            product.save()
            updated += 1
        else:
            new_products.append(Product(
                name=result['name'],
                slug=_unique_slug(result['name'], taken_slugs),
                company=companies[result['company_id']],
                company_line=lines.get(result['line_id']) if result['line_id'] else None,
                category=categories[result['category_id']],
                description=result['description'],
                size=result['sizes'],
                is_available=result['available'],
                meta_title=result['seo_title'] or None,
                meta_description=result['seo_description'] or None,
                youtube_url=result['youtube'],
            ))

    # New products have no image yet, so save() (which converts images) has
    # nothing to do; inserting in batches keeps big files fast on the server
    Product.objects.bulk_create(new_products, batch_size=500)
    return len(new_products), updated
